"""
Servicio para generación de reportes Excel.

Usa openpyxl para generar archivos .xlsx en memoria.

Instalación:
pip install openpyxl
"""

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1E3A8A"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)


def normalize_excel_value(value: Any) -> Any:
    """
    Normaliza valores antes de escribirlos en Excel.

    openpyxl no acepta bien datetime con timezone,
    por eso se convierten a texto ISO.
    """
    if isinstance(value, datetime):
        return value.isoformat()

    return value


def style_header_row(sheet) -> None:
    """
    Aplica estilo a la primera fila de una hoja.
    """
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def autosize_columns(sheet) -> None:
    """
    Ajusta ancho de columnas de forma simple.
    """
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            max_length = max(
                max_length,
                len(str(value))
            )

        sheet.column_dimensions[column_letter].width = max_length + 2


def workbook_to_bytes(workbook: Workbook) -> bytes:
    """
    Convierte workbook a bytes.
    """
    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()


def build_sales_excel(data: dict[str, Any]) -> bytes:
    """
    Genera Excel de ventas.
    """
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "KPI"

    summary_sheet.append([
        "Indicador",
        "Valor"
    ])

    kpis = data.get("kpis", {})

    summary_sheet.append(["Ventas totales", kpis.get("totalSales", 0)])
    summary_sheet.append(["Pedidos totales", kpis.get("totalOrders", 0)])
    summary_sheet.append(["Productos vendidos", kpis.get("totalProducts", 0)])
    summary_sheet.append(["Clientes totales", kpis.get("totalCustomers", 0)])
    summary_sheet.append(["Ticket promedio", kpis.get("averageTicket", 0)])
    summary_sheet.append(["Pedidos pendientes", kpis.get("pendingOrders", 0)])
    summary_sheet.append(["Pedidos pagados", kpis.get("paidOrders", 0)])
    summary_sheet.append(["Pedidos cancelados", kpis.get("cancelledOrders", 0)])
    summary_sheet.append(["Pedidos entregados", kpis.get("deliveredOrders", 0)])

    sales_sheet = workbook.create_sheet("Ventas por fecha")

    sales_sheet.append([
        "Fecha",
        "Ventas totales",
        "Pedidos"
    ])

    for item in data.get("salesSummary", []):
        sales_sheet.append([
            normalize_excel_value(item.get("date")),
            item.get("totalSales", 0),
            item.get("totalOrders", 0)
        ])

    products_sheet = workbook.create_sheet("Productos más vendidos")

    products_sheet.append([
        "Producto ID",
        "Producto",
        "Categoría",
        "Cantidad",
        "Ventas"
    ])

    for item in data.get("topProducts", []):
        products_sheet.append([
            item.get("productId"),
            item.get("productName"),
            item.get("category"),
            item.get("totalQuantity", 0),
            item.get("totalSales", 0)
        ])

    categories_sheet = workbook.create_sheet("Ventas por categoría")

    categories_sheet.append([
        "Categoría",
        "Ventas",
        "Cantidad"
    ])

    for item in data.get("salesByCategory", []):
        categories_sheet.append([
            item.get("category"),
            item.get("totalSales", 0),
            item.get("totalQuantity", 0)
        ])

    for sheet in workbook.worksheets:
        style_header_row(sheet)
        autosize_columns(sheet)

    return workbook_to_bytes(workbook)


def build_products_excel(data: dict[str, Any]) -> bytes:
    """
    Genera Excel de productos.
    """
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Productos"

    sheet.append([
        "ID",
        "Nombre",
        "Descripción",
        "Categoría",
        "Precio",
        "Stock",
        "Estado",
        "Destacado",
        "Fecha de creación"
    ])

    for product in data.get("products", []):
        sheet.append([
            product.get("id"),
            product.get("name"),
            product.get("description"),
            product.get("category"),
            product.get("price"),
            product.get("stock"),
            product.get("status"),
            "Sí" if product.get("featured") else "No",
            normalize_excel_value(product.get("createdAt"))
        ])

    style_header_row(sheet)
    autosize_columns(sheet)

    return workbook_to_bytes(workbook)


def build_orders_excel(data: dict[str, Any]) -> bytes:
    """
    Genera Excel de pedidos.
    """
    workbook = Workbook()

    orders_sheet = workbook.active
    orders_sheet.title = "Pedidos"

    orders_sheet.append([
        "ID",
        "Folio",
        "Cliente",
        "Estado",
        "Método de pago",
        "Total de productos",
        "Subtotal",
        "Impuesto",
        "Descuento",
        "Total",
        "Fecha"
    ])

    details_sheet = workbook.create_sheet("Detalles")

    details_sheet.append([
        "Folio",
        "Producto ID",
        "Producto",
        "Categoría",
        "Precio unitario",
        "Cantidad",
        "Subtotal"
    ])

    for order in data.get("orders", []):
        orders_sheet.append([
            order.get("id"),
            order.get("folio"),
            order.get("customerName"),
            order.get("status"),
            order.get("paymentMethod"),
            order.get("totalItems"),
            order.get("subtotal"),
            order.get("tax"),
            order.get("discount"),
            order.get("total"),
            normalize_excel_value(order.get("createdAt"))
        ])

        for detail in order.get("details", []):
            details_sheet.append([
                order.get("folio"),
                detail.get("productId"),
                detail.get("productName"),
                detail.get("category"),
                detail.get("unitPrice"),
                detail.get("quantity"),
                detail.get("subtotal")
            ])

    for sheet in workbook.worksheets:
        style_header_row(sheet)
        autosize_columns(sheet)

    return workbook_to_bytes(workbook)